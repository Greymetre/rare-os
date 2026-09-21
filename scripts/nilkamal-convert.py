#!/usr/bin/env python3
"""Converts the Nilkamal simulation handover (21-Sep-2026) into one load bundle for
scripts/load-nilkamal.mjs. Barjora plant (1116) only, reproducing the frozen demo inputs:

  items / BOMs / buffers / production orders   Demo_Build/nilkamal_seed_v9_materials.json
  component stock (unrestricted, per SLoc)      1116-Barjora.xlsx, sheet MB52
  finished-goods stock                         the demo's FG buffer positions (not MB52)
  open purchase order lines (balance qty)      R1 A5 -open purchase orders.xlsx, sheet 1116
  daily demand history (net of returns)        Sleep Sale From Apr 24 to Jun 26.xlsx

The bundle is client data: it is written under .local/ (git- and docker-ignored) and must
never be committed. Needs Python 3 with openpyxl.

Usage: python3 scripts/nilkamal-convert.py <handover folder> [.local/nilkamal/bundle.json]
"""
import collections
import datetime
import hashlib
import json
import math
import os
import re
import sys

import openpyxl

HANDOVER = sys.argv[1] if len(sys.argv) > 1 else '../../Nilkamal_Simulation_Demo_Developer_Handover_21Sep2026'
OUT = sys.argv[2] if len(sys.argv) > 2 else '.local/nilkamal/bundle.json'
DATA = os.path.join(HANDOVER, 'Simulation data')
MODEL_DATE = '2026-07-27'
HISTORY_END = datetime.date(2026, 6, 30)
PLANT = '1116'
LEAD_TIME = 10  # the demo's supplier planning assumption; A1 purchase lead times were not supplied
NO_SOURCE = 'NK-NOSOURCE'


def md5(path):
    with open(path, 'rb') as f:
        return hashlib.md5(f.read()).hexdigest()


def text(v, n=120):
    s = re.sub(r'\s+', ' ', str(v or '')).strip()
    return s[:n]


def unit(u):
    u = str(u or '').strip().upper()
    return {'KGS': 'KG'}.get(u, u)


def iso(d):
    if isinstance(d, datetime.datetime):
        return d.date().isoformat()
    m = re.match(r'^(\d{1,2})-([A-Za-z]{3})-(\d{4})$', str(d))
    if m:
        return datetime.datetime.strptime(str(d), '%d-%b-%Y').date().isoformat()
    m = re.match(r'^(\d{2})\.(\d{2})\.(\d{4})$', str(d))
    if m:
        return f'{m[3]}-{m[2]}-{m[1]}'
    raise ValueError(f'Unknown date {d!r}')


seed_path = os.path.join(HANDOVER, 'Demo_Build', 'nilkamal_seed_v9_materials.json')
seed = json.load(open(seed_path))
fg_rows = seed['finished_goods']
fgs = [f['code'] for f in fg_rows]
positions = seed['material_positions']
bom_lines = {f: seed['bom_lines'][f] for f in fgs}
components = []
for f in fgs:
    for c, *_ in bom_lines[f]:
        if c not in components:
            components.append(c)
current = set(fgs) | set(components)

# ---------- Demand history: national invoice quantity per FG and day, returns netted ----------
sales_path = os.path.join(DATA, 'Nilkamal Mattress', 'Sleep Sale From Apr 24 to Jun 26.xlsx')
daily = collections.defaultdict(float)
ws = openpyxl.load_workbook(sales_path, read_only=True).worksheets[0]
for i, r in enumerate(ws.iter_rows(values_only=True)):
    if i == 0:
        continue
    d, material, qty = r[11], r[21], r[23]
    if not isinstance(d, datetime.datetime) or material is None or qty is None:
        continue  # one corrupt source line
    material = str(material).strip()
    if material in fgs:
        daily[(material, d.date().isoformat())] += qty
demand = [[m, d, round(q, 6)] for (m, d), q in sorted(daily.items()) if q != 0]

# Component ADU for the MOQ planning parameter: trailing 91 days of parent demand x full BOM.
def adu91(code):
    return sum(q for (m, d), q in daily.items()
               if m == code and (HISTORY_END - datetime.date.fromisoformat(d)).days < 91) / 91

fg_adu = {f: adu91(f) for f in fgs}
comp_adu = collections.defaultdict(float)
for f in fgs:
    for c, q, *_ in bom_lines[f]:
        comp_adu[c] += fg_adu[f] * q

# ---------- Stock: MB52 unrestricted per storage location ----------
barjora_path = os.path.join(DATA, 'Nilkamal Mattress', '1116-Barjora.xlsx')
stock = collections.defaultdict(float)
stock_units = {}
ws = openpyxl.load_workbook(barjora_path, read_only=True)['MB52']
for i, r in enumerate(ws.iter_rows(values_only=True)):
    if i == 0 or r[0] is None:
        continue
    material = str(r[0]).strip()
    # Finished goods use the demo's FG buffer positions, not MB52 (see the reconciliation below).
    if material not in components or str(r[2]).strip() != PLANT:
        continue
    stock[(material, str(r[3]).strip() if r[3] else 'NOSLOC')] += float(r[5] or 0)
    stock_units[material] = unit(r[4])

# ---------- Open purchase orders: A5 plant sheet, balance quantity ----------
po_path = os.path.join(DATA, 'Nilkamal Mattress R1_06Aug2026', 'A5 -open purchase orders.xlsx')
wb = openpyxl.load_workbook(po_path, read_only=True)
vendors = {}
for sheet in wb.worksheets:
    for i, r in enumerate(sheet.iter_rows(values_only=True)):
        if i and r[10] and r[11]:
            vendors.setdefault(text(r[11]).upper(), (str(r[10]).strip(), text(r[11])))
po_lines = []
ws = wb[PLANT]
for i, r in enumerate(ws.iter_rows(values_only=True)):
    if i == 0 or not r[4]:
        continue
    material = str(r[4]).strip()
    if material not in components or not r[8] or float(r[8]) <= 0:
        continue
    po_lines.append({
        'po_no': str(r[2]).strip(), 'line_no': int(str(r[52]).strip()), 'item': material,
        'quantity': float(r[8]), 'unit': unit(r[124]), 'due_date': iso(r[35]),
        'order_date': iso(r[3]), 'supplier': str(r[10]).strip(), 'row': i + 1,
    })

# SAP schedule lines repeat a PO item with different delivery dates; each becomes its own line
# numbered item x 10 + schedule (80 -> 801, 802, ...), keeping quantity and date.
schedules = collections.Counter((l['po_no'], l['line_no']) for l in po_lines)
seen_schedule = collections.Counter()
for l in po_lines:
    k = (l['po_no'], l['line_no'])
    if schedules[k] > 1:
        seen_schedule[k] += 1
        l['sap_item'] = l['line_no']
        l['line_no'] = l['line_no'] * 10 + seen_schedule[k]

# ---------- Masters ----------
def decimals_of(values):
    best = 0
    for v in values:
        s = repr(round(float(v), 6))
        if '.' in s and not s.endswith('.0'):
            best = max(best, len(s.split('.')[1]))
    return min(6, best)

by_unit = collections.defaultdict(list)
for f in fgs:
    for c, q, u, _ in bom_lines[f]:
        by_unit[unit(u)].append(q)
for (m, _), q in stock.items():
    by_unit[stock_units[m]].append(q)
for l in po_lines:
    by_unit[l['unit']].append(l['quantity'])
UNIT_NAMES = {'NOS': 'Numbers', 'KG': 'Kilogram', 'M': 'Metre', 'L': 'Litre'}
units = [[u, UNIT_NAMES.get(u, u), max(0, decimals_of(v))] for u, v in sorted(by_unit.items())]
units = [u for u in units if u[0] in UNIT_NAMES] + [u for u in units if u[0] not in UNIT_NAMES]

CLASS = {'runner': 'runner', 'repeater': 'repeater', 'stranger': 'stranger'}
items = [{'code': f['code'], 'name': text(f['desc']) or f['code'], 'type': 'FG', 'make_buy': 'MAKE',
          'unit': 'NOS', 'family': text(f.get('family'), 60), 'demand_class': CLASS.get(f.get('demand_class'))}
         for f in fg_rows]
for c in components:
    p = positions[c]
    items.append({'code': c, 'name': text(p.get('name')) or c, 'type': 'RM', 'make_buy': 'BUY',
                  'unit': unit(p['uom']), 'family': '', 'demand_class': None})

# Preferred source: the first supplier the demo names; items without an open PO get a placeholder.
suppliers = {NO_SOURCE: 'Supplier not in the open-PO extract'}
sources = []
for c in components:
    names = [n.strip() for n in str(positions[c].get('supplier') or '').split(';') if n.strip()]
    code = NO_SOURCE
    for n in names:
        if n.upper() in vendors:
            code = vendors[n.upper()][0]
            suppliers[code] = vendors[n.upper()][1]
            break
    moq = max(10, math.floor(comp_adu[c] * 1.5 / 10 + 0.5) * 10)  # the demo's MOQ rule (JS rounding)
    sources.append({'item': c, 'supplier': code, 'moq': moq, 'multiple': 50})
for l in po_lines:
    suppliers.setdefault(l['supplier'], next((v[1] for v in vendors.values() if v[0] == l['supplier']), l['supplier']))

# Work centres and routings (the frozen seed's Barjora 16-operation book).
ops = seed['routing_ops']
op_names = {o['op']: o['name'] for o in seed['op_master']}
op_seq = {o['op']: o['seq'] for o in seed['op_master']}
resources = []
for w in seed['work_centres']:
    if w['site'] != 'PLANT-' + PLANT:
        continue
    resources.append({'code': w['op'], 'name': text(w['name']), 'machines': max(1, int(round(w['machine_count']))),
                      'efficiency': w['efficiency_pct'], 'changeover': w['changeover_min']})
routings = {}
for f in fgs:
    mins = seed['routing_min_per_unit'].get(f) or []
    routings[f] = [[op_seq[op], op, op_names[op], m] for op, m in zip(ops, mins) if m and m > 0]

buffers = [{'item': b['item'], 'lead_time_days': b['lt_days'] if not b.get('component') else None,
            'kind': 'RM' if b.get('component') else 'FG'} for b in seed['buffers']]
buffered = {b['item'] for b in buffers}

production = [{'order_no': o['id'], 'item': o['fg'], 'quantity': o['qty'], 'start_date': iso(o['bsc_start']),
               'due_date': iso(o['promise_date']), 'order_type': o.get('order_type', ''),
               'reference': 'A5 open production orders, sheet 1116'}
              for o in seed['sales_orders'] if not re.search('SPIKE', o.get('status', ''), re.I)]

# ---------- Reconcile with the demo before writing ----------
problems = []
on_hand = collections.defaultdict(float)
for (m, _), q in stock.items():
    on_hand[m] += q
for c in components:
    p = positions[c]
    if p.get('stock_present') and abs(on_hand.get(c, 0) - p['on_hand']) > 1e-6:
        problems.append(f'stock {c}: {on_hand.get(c)} vs demo {p["on_hand"]}')
    if not p.get('stock_present') and c in on_hand:
        problems.append(f'stock {c}: MB52 row but the demo has none')
    oo = sum(l['quantity'] for l in po_lines if l['item'] == c)
    if abs(oo - (p.get('on_order') or 0)) > 1e-6:
        problems.append(f'on order {c}: {oo} vs demo {p.get("on_order")}')
    if c in stock_units and stock_units[c] != unit(p['uom']):
        problems.append(f'unit {c}: stock {stock_units[c]} vs BOM {p["uom"]}')
FG_STORE = 'FG01'
for b in seed['buffers']:
    if not b.get('component') and b['on_hand'] > 0:
        stock[(b['item'], FG_STORE)] += b['on_hand']
if problems:
    sys.exit('Reconciliation failed:\n  ' + '\n  '.join(problems[:40]))

bundle = {
    'meta': {
        'company': 'Nilkamal', 'org': seed['meta']['org'], 'model_date': MODEL_DATE,
        'history_end': HISTORY_END.isoformat(), 'lead_time_days': LEAD_TIME,
        'plant': {'code': PLANT, 'name': 'Barjora Mattress Mfg. Unit', 'location': 'Barjora, West Bengal'},
        'sources': {'seed': md5(seed_path), 'sales': md5(sales_path), 'barjora': md5(barjora_path), 'open_po': md5(po_path)},
        'note': 'Nilkamal client data for development and simulation only. Do not commit or share.',
    },
    'units': units,
    'items': items,
    'suppliers': [[k, v] for k, v in sorted(suppliers.items())],
    'sources': sources,
    'locations': sorted({l for (_, l) in stock}),
    'fg_location': FG_STORE,
    'stock': [[m, l, round(q, 6)] for (m, l), q in sorted(stock.items()) if q > 0],
    'purchase_orders': po_lines,
    'production_orders': production,
    'demand': demand,
    'boms': {f: [[c, q, unit(u)] for c, q, u, _ in bom_lines[f]] for f in fgs},
    'resources': resources,
    'routings': routings,
    'buffers': buffers,
    'mto': [f for f in fgs if f not in buffered],
}
os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, 'w') as f:
    json.dump(bundle, f)
print(f'{OUT}: {len(items)} items ({len(fgs)} FG, {len(components)} components), {len(bundle["stock"])} stock rows '
      f'in {len(bundle["locations"])} locations, {len(po_lines)} open PO lines, {len(production)} production orders, '
      f'{len(demand)} demand days, {len(buffers)} buffers, {len(resources)} work centres. Reconciled with the demo.')
